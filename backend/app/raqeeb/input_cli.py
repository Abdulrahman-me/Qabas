"""Isolated, bounded attachment decoding. Binary input/output stays off queues and logs."""
from __future__ import annotations

import base64
import importlib
import io
import json
import os
import subprocess
import sys
import zipfile
from typing import Any


def constrain(megabytes: int) -> Any:
    if sys.platform != "win32":
        resource = importlib.import_module("resource")
        resource.setrlimit(resource.RLIMIT_AS, (megabytes * 1024**2, megabytes * 1024**2))
        return None
    import ctypes
    from ctypes import wintypes

    class Basic(ctypes.Structure):
        _fields_ = [("PerProcessUserTimeLimit", ctypes.c_int64), ("PerJobUserTimeLimit", ctypes.c_int64),
            ("LimitFlags", wintypes.DWORD), ("MinimumWorkingSetSize", ctypes.c_size_t),
            ("MaximumWorkingSetSize", ctypes.c_size_t), ("ActiveProcessLimit", wintypes.DWORD),
            ("Affinity", ctypes.c_size_t), ("PriorityClass", wintypes.DWORD), ("SchedulingClass", wintypes.DWORD)]

    class Counters(ctypes.Structure):
        _fields_ = [(n, ctypes.c_uint64) for n in ("ReadOperationCount", "WriteOperationCount", "OtherOperationCount",
                                                "ReadTransferCount", "WriteTransferCount", "OtherTransferCount")]

    class Extended(ctypes.Structure):
        _fields_ = [("BasicLimitInformation", Basic), ("IoInfo", Counters),
                    ("ProcessMemoryLimit", ctypes.c_size_t), ("JobMemoryLimit", ctypes.c_size_t),
                    ("PeakProcessMemoryUsed", ctypes.c_size_t), ("PeakJobMemoryUsed", ctypes.c_size_t)]

    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.CreateJobObjectW.restype = wintypes.HANDLE
    kernel.GetCurrentProcess.restype = wintypes.HANDLE
    kernel.SetInformationJobObject.argtypes = [wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p, wintypes.DWORD]
    kernel.AssignProcessToJobObject.argtypes = [wintypes.HANDLE, wintypes.HANDLE]
    handle = kernel.CreateJobObjectW(None, None)
    limits = Extended()
    limits.BasicLimitInformation.LimitFlags = 0x2000 | 0x200  # kill on handle close + total job memory
    limits.JobMemoryLimit = megabytes * 1024**2
    if not handle or not kernel.SetInformationJobObject(handle, 9, ctypes.byref(limits), ctypes.sizeof(limits)) or \
            not kernel.AssignProcessToJobObject(handle, kernel.GetCurrentProcess()):
        raise OSError("cannot enforce input memory limit")
    return handle  # keep the handle alive until process exit, including ffmpeg children


def image_bytes(image: Any, *, png: bool = False) -> bytes:
    image = image.convert("RGB")
    image.thumbnail((1800, 1800))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG" if png else "JPEG", **({} if png else {"quality": 85}))
    return buffer.getvalue()


def decode(value: dict[str, Any]) -> dict[str, Any]:
    from PIL import Image

    data, mime = base64.b64decode(value["data"], validate=True), value["mime"]
    if value["kind"] == "image":
        with Image.open(io.BytesIO(data)) as original:
            if Image.MIME.get(original.format or "") != mime or original.width * original.height > 40_000_000:
                raise ValueError("image signature/dimensions")
            original.verify()
        with Image.open(io.BytesIO(data)) as original:
            rendered = image_bytes(original)
        return {"images": [base64.b64encode(rendered).decode()], "text": "", "pages": None,
                "pages_processed": 0, "truncated": False, "duration_ms": None, "wav": None}
    if value["kind"] == "audio":
        import av

        signatures = {"audio/wav": data.startswith(b"RIFF") and data[8:12] == b"WAVE",
            "audio/mp4": data[4:8] == b"ftyp", "audio/webm": data.startswith(b"\x1aE\xdf\xa3"),
            "audio/mpeg": data.startswith(b"ID3") or (len(data) >= 2 and data[0] == 255 and data[1] & 224 == 224)}
        if not signatures.get(mime):
            raise ValueError("audio signature")
        with av.open(io.BytesIO(data), options={"protocol_whitelist": "pipe,data", "enable_drefs": "0"}) as source:
            if not isinstance(source, av.container.InputContainer) or not source.streams.audio:
                raise ValueError("missing audio stream")
            codec = source.streams.audio[0].codec_context.name
            if mime in {"audio/mp4", "audio/webm", "audio/mpeg"} and codec not in {
                    "audio/mp4": {"aac", "aac_fixed"}, "audio/webm": {"opus"},
                    "audio/mpeg": {"mp3", "mp3float"}}[mime]:
                raise ValueError("unsupported audio codec")
        result = subprocess.run([value["ffmpeg"], "-nostdin", "-hide_banner", "-loglevel", "error",  # noqa: S603
            "-protocol_whitelist", "pipe,data", "-i", "pipe:0",
            "-t", "61", "-vn", "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le", "-f", "wav", "pipe:1"],
            input=data, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=8, check=True)
        with av.open(io.BytesIO(result.stdout)) as container:
            assert isinstance(container, av.container.InputContainer)
            samples = 0
            for frame in container.decode(audio=0):
                assert isinstance(frame, av.AudioFrame)
                samples += frame.samples
        if not 0 < samples <= 960_000:
            raise ValueError("audio duration")
        return {"images": [], "text": "", "pages": None, "pages_processed": 0, "truncated": False,
                "duration_ms": round(samples * 1000 / 16000), "wav": base64.b64encode(result.stdout).decode()}
    images, pages, processed = [], None, 0
    if mime == "application/pdf":
        import pdfplumber

        if not data.startswith(b"%PDF-"):
            raise ValueError("PDF signature")
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if pdf.doc.encryption:
                raise ValueError("encrypted PDF")
            pages = len(pdf.pages)
            if not pages:
                raise ValueError("empty PDF")
            processed = min(20, pages)
            text = "\n".join(p.extract_text() or "" for p in pdf.pages[:20])
        if len(text.strip()) < 200:
            import pypdfium2 as pdfium

            with pdfium.PdfDocument(data) as pdf:
                processed = min(5, len(pdf))
                for index in range(processed):
                    page = pdf[index]
                    width, height = page.get_size()
                    bitmap = page.render(scale=min(2, 1800 / max(width, height)))
                    try:
                        image = bitmap.to_pil()
                        images.append(base64.b64encode(image_bytes(image)).decode())
                    finally:
                        bitmap.close()
                        page.close()
    else:
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 50 * 1024**2 or any(
                    e.flag_bits & 1 or e.filename.lower().endswith("vbaproject.bin") for e in entries):
                raise ValueError("encrypted/macro/expanded DOCX")
            content_types = archive.read("[Content_Types].xml")
            if b"macroenabled" in content_types.lower() or b"vbaproject" in content_types.lower() or \
                    "word/document.xml" not in archive.namelist():
                raise ValueError("not plain DOCX")
        document = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
        # Paragraph extraction does not render Word pagination. Do not invent a page count.
    if not text.strip() and not images:
        raise ValueError("no readable document content")
    return {"images": images, "text": text[:30_000], "pages": pages, "pages_processed": processed,
            "truncated": len(text) > 30_000 or (pages is not None and pages > processed),
            "duration_ms": None, "wav": None}


def main() -> int:
    try:
        value = json.loads(sys.stdin.buffer.read(15_000_000))
        job_handle = constrain(value["memory_mb"])
        result = decode(value)
        result["memory_enforced"] = job_handle is not None or os.name != "nt"
        output = json.dumps(result).encode()
        if len(output) > 8_000_000:
            raise ValueError("decoded input output exceeds limit")
        sys.stdout.buffer.write(output)
        return 0
    except (FileNotFoundError, ImportError):
        return 3  # unavailable native capability, distinct from a corrupt file
    except Exception:
        return 2  # no filename, text, parser diagnostics or media bytes on stderr/logs


if __name__ == "__main__":
    raise SystemExit(main())
